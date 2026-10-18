import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
        else:
            tag = f'w:{edge}'
            element = OxmlElement(tag)
            element.set(qn('w:val'), 'none')
            tcBorders.append(element)
    tcPr.append(tcBorders)

def build_salary_increase_doc():
    doc = Document()

    # Brand Colors
    COLOR_VERDE = RGBColor(83, 88, 57)      # #535839 (Verde Oliva oficial)
    COLOR_CAFE = RGBColor(38, 18, 15)       # #26120f (Café Oscuro oficial)
    COLOR_DORADO = RGBColor(208, 187, 125)  # #d0bb7d (Dorado oficial)
    COLOR_GRIS = RGBColor(121, 135, 155)    # #79879b (Azul Gris Medio)
    COLOR_MUTED = RGBColor(100, 116, 139)   # Slate gray

    # Configure Margins for Letter Page (8.5 x 11 in) - Calibrados para ajuste impecable
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # ----------------------------------------------------
    # 1. HEADER (Logo on left, Brand & Legal info on right)
    # ----------------------------------------------------
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    row = header_table.rows[0]
    cell_logo = row.cells[0]
    cell_text = row.cells[1]
    cell_logo.width = Inches(2.3)
    cell_text.width = Inches(4.7)

    logo_path = os.path.abspath(r'assets/images/logos/logo-header-verdeoliva.png')
    p_logo = cell_logo.paragraphs[0]
    p_logo.paragraph_format.space_after = Pt(0)
    p_logo.paragraph_format.space_before = Pt(0)
    r_logo = p_logo.add_run()
    if os.path.exists(logo_path):
        r_logo.add_picture(logo_path, width=Inches(1.95))

    p_h1 = cell_text.paragraphs[0]
    p_h1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h1.paragraph_format.space_after = Pt(1)
    p_h1.paragraph_format.space_before = Pt(2)
    r_h1 = p_h1.add_run("BAKE ME UP")
    r_h1.font.name = 'Segoe UI'
    r_h1.font.size = Pt(12.5)
    r_h1.font.bold = True
    r_h1.font.color.rgb = COLOR_VERDE

    p_h2 = cell_text.add_paragraph()
    p_h2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h2.paragraph_format.space_after = Pt(1)
    r_h2 = p_h2.add_run("Repostería Fina & Café de Especialidad")
    r_h2.font.name = 'Segoe UI'
    r_h2.font.size = Pt(8.5)
    r_h2.font.italic = True
    r_h2.font.color.rgb = COLOR_GRIS

    p_h3 = cell_text.add_paragraph()
    p_h3.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h3.paragraph_format.space_after = Pt(0)
    r_h3 = p_h3.add_run("RFC: GAAS0208311X9 · Régimen 626 (RESICO)\nPuebla, Pue. · bakemeup.mx")
    r_h3.font.name = 'Segoe UI'
    r_h3.font.size = Pt(7.5)
    r_h3.font.color.rgb = COLOR_MUTED

    # Pleca divisoria
    divider_table = doc.add_table(rows=1, cols=1)
    divider_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_div = divider_table.rows[0].cells[0]
    c_div.width = Inches(7.0)
    set_cell_background(c_div, "535839")
    set_cell_margins(c_div, top=12, bottom=12, left=0, right=0)
    p_div = c_div.paragraphs[0]
    p_div.paragraph_format.space_after = Pt(0)
    p_div.paragraph_format.space_before = Pt(0)
    r_div = p_div.add_run("")
    r_div.font.size = Pt(1)

    # ----------------------------------------------------
    # 2. METADATA / DATE & ADDRESSEE BLOCK
    # ----------------------------------------------------
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(8)
    p_meta.paragraph_format.space_after = Pt(6)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    r_loc = p_meta.add_run("Puebla, Pue., a 14 de octubre de 2026\n")
    r_loc.font.name = 'Segoe UI'
    r_loc.font.size = Pt(9)
    r_loc.font.bold = True
    r_loc.font.color.rgb = COLOR_CAFE

    r_asunto = p_meta.add_run("ASUNTO: ASIGNACIÓN DE NUEVAS RESPONSABILIDADES E INCREMENTO SALARIAL")
    r_asunto.font.name = 'Segoe UI'
    r_asunto.font.size = Pt(8)
    r_asunto.font.bold = True
    r_asunto.font.color.rgb = COLOR_VERDE

    # Destinatario: JOSÉ DEL CASTILLO OLVERA (Encargado de Producción)
    p_dest = doc.add_paragraph()
    p_dest.paragraph_format.space_before = Pt(0)
    p_dest.paragraph_format.space_after = Pt(6)
    p_dest.paragraph_format.line_spacing = 1.10

    r_dest_title = p_dest.add_run("C. JOSÉ DEL CASTILLO OLVERA\n")
    r_dest_title.font.name = 'Segoe UI'
    r_dest_title.font.size = Pt(10)
    r_dest_title.font.bold = True
    r_dest_title.font.color.rgb = COLOR_CAFE

    r_dest_cargo = p_dest.add_run("Puesto: Encargado de Producción\nÁrea: Taller y Cocina de Producción · Bake Me Up\nP R E S E N T E .")
    r_dest_cargo.font.name = 'Segoe UI'
    r_dest_cargo.font.size = Pt(9)
    r_dest_cargo.font.color.rgb = COLOR_CAFE

    # ----------------------------------------------------
    # 3. BODY TEXT (Reconocimiento y Promoción de Funciones)
    # ----------------------------------------------------
    p_b1 = doc.add_paragraph()
    p_b1.paragraph_format.space_before = Pt(2)
    p_b1.paragraph_format.space_after = Pt(5)
    p_b1.paragraph_format.line_spacing = 1.10

    r_b1 = p_b1.add_run(
        "Por medio de la presente, la Dirección General de Bake Me Up reconoce tu constancia, compromiso y calidad demostrada "
        "en el taller de producción. Con motivo de tu crecimiento dentro del equipo y en respuesta a la confianza que has consolidado, "
        "se formaliza tu promoción al puesto de "
    )
    r_b1.font.name = 'Segoe UI'
    r_b1.font.size = Pt(8.5)
    r_b1.font.color.rgb = COLOR_CAFE

    r_b1_puesto = p_b1.add_run("Encargado de Producción")
    r_b1_puesto.font.name = 'Segoe UI'
    r_b1_puesto.font.size = Pt(8.5)
    r_b1_puesto.font.bold = True
    r_b1_puesto.font.color.rgb = COLOR_VERDE

    r_b1_end = p_b1.add_run(", asumiendo a partir de esta fecha la titularidad y liderazgo de las siguientes responsabilidades operativas:")
    r_b1_end.font.name = 'Segoe UI'
    r_b1_end.font.size = Pt(8.5)
    r_b1_end.font.color.rgb = COLOR_CAFE

    # ----------------------------------------------------
    # 4. TABLA DE RESPONSABILIDADES Y ALCANCE OPERATIVO
    # ----------------------------------------------------
    resp_table = doc.add_table(rows=3, cols=2)
    resp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    resp_table.autofit = False

    responsabilidades = [
        ("1. Abastecimiento y Proveedores Fijos",
         "Revisión continua de inventario de materias primas y gestión directa de pedidos ante proveedores establecidos de abarrotes y lácteos para asegurar abasto oportuno en cocina (la Dirección autoriza y efectúa pagos)."),
        ("2. Autonomía y Control de Producción",
         "Monitoreo proactivo de existencias en taller, detectando faltantes de bases, batidos y producto terminado, asumiendo su elaboración de forma autónoma sin requerir instrucción puntual de la Dirección."),
        ("3. Planificación de Pedidos y Horneado",
         "Revisión anticipada diaria de la agenda de pedidos programados de pasteles y eventos, coordinando con antelación los tiempos óptimos de horneado, enfriamiento, preparación y montaje.")
    ]

    for idx, (titulo, desc) in enumerate(responsabilidades):
        r_row = resp_table.rows[idx]
        c0, c1 = r_row.cells[0], r_row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)
        
        set_cell_background(c0, "F8F9F6")
        set_cell_background(c1, "F8F9F6")
        set_cell_margins(c0, top=35, bottom=35, left=70, right=50)
        set_cell_margins(c1, top=35, bottom=35, left=50, right=70)

        b_color = "535839"
        border_cfg = {
            'top': {'val': 'single', 'sz': 4, 'color': b_color} if idx == 0 else {'val': 'dotted', 'sz': 3, 'color': 'D0BB7D'},
            'bottom': {'val': 'single', 'sz': 4, 'color': b_color} if idx == 2 else {'val': 'none'},
            'left': {'val': 'single', 'sz': 4, 'color': b_color},
            'right': {'val': 'single', 'sz': 4, 'color': b_color}
        }
        set_cell_border(c0, **border_cfg)
        set_cell_border(c1, **border_cfg)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.line_spacing = 1.05
        r0 = p0.add_run(titulo)
        r0.font.name = 'Segoe UI'
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_VERDE

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.line_spacing = 1.05
        r1 = p1.add_run(desc)
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(8)
        r1.font.color.rgb = COLOR_CAFE

    # Párrafo de enlace hacia la compensación
    p_b2 = doc.add_paragraph()
    p_b2.paragraph_format.space_before = Pt(5)
    p_b2.paragraph_format.space_after = Pt(5)
    p_b2.paragraph_format.line_spacing = 1.10

    r_b2 = p_b2.add_run(
        "En correlación directa a la ampliación de estas responsabilidades y el nivel de liderazgo asignado, "
        "la Dirección aprueba una mejora en tu esquema salarial mensual sobre la compensación previa de "
    )
    r_b2.font.name = 'Segoe UI'
    r_b2.font.size = Pt(8.5)
    r_b2.font.color.rgb = COLOR_CAFE

    r_b2_bold = p_b2.add_run("$8,000.00 MXN mensuales")
    r_b2_bold.font.name = 'Segoe UI'
    r_b2_bold.font.size = Pt(8.5)
    r_b2_bold.font.bold = True
    r_b2_bold.font.color.rgb = COLOR_VERDE

    r_b2_end = p_b2.add_run(", formalizándose el nuevo monto acordado en los siguientes términos:")
    r_b2_end.font.name = 'Segoe UI'
    r_b2_end.font.size = Pt(8.5)
    r_b2_end.font.color.rgb = COLOR_CAFE

    # ----------------------------------------------------
    # 5. HIGHLIGHT BOX (TABLA DE SUELDO CON LÍNEA EN BLANCO PARA PLUMA)
    # ----------------------------------------------------
    box_table = doc.add_table(rows=4, cols=2)
    box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    box_table.autofit = False

    labels_data = [
        ("Salario Mensual Previo:", "$8,000.00 MXN (Ocho mil pesos 00/100 M.N.)"),
        ("NUEVO SALARIO MENSUAL ASIGNADO:", "$ __________________________________ MXN"),
        ("Importe del Nuevo Salario en Letra:", "( __________________________________________________________________ )"),
        ("Fecha de Entrada en Vigor:", "14 de octubre de 2026")
    ]

    for idx, (label, val) in enumerate(labels_data):
        row = box_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.6)
        c1.width = Inches(4.4)
        
        set_cell_background(c0, "F8F9F6")
        set_cell_background(c1, "F8F9F6")
        set_cell_margins(c0, top=40, bottom=40, left=80, right=60)
        set_cell_margins(c1, top=40, bottom=40, left=60, right=80)

        b_color = "535839"
        border_cfg = {
            'top': {'val': 'single', 'sz': 4, 'color': b_color} if idx == 0 else {'val': 'dotted', 'sz': 3, 'color': 'D0BB7D'},
            'bottom': {'val': 'single', 'sz': 4, 'color': b_color} if idx == 3 else {'val': 'none'},
            'left': {'val': 'single', 'sz': 4, 'color': b_color},
            'right': {'val': 'single', 'sz': 4, 'color': b_color}
        }
        set_cell_border(c0, **border_cfg)
        set_cell_border(c1, **border_cfg)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.space_before = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = 'Segoe UI'
        r0.font.size = Pt(8)
        if "NUEVO SALARIO" in label:
            r0.font.bold = True
            r0.font.color.rgb = COLOR_VERDE
        else:
            r0.font.bold = True
            r0.font.color.rgb = COLOR_CAFE

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_before = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(8)
        if "NUEVO SALARIO" in label:
            r1.font.bold = True
            r1.font.color.rgb = COLOR_CAFE
        elif "Letra" in label:
            r1.font.italic = True
            r1.font.color.rgb = COLOR_MUTED
        else:
            r1.font.color.rgb = COLOR_CAFE

    # ----------------------------------------------------
    # 6. SIGNATURE BLOCK (2 Columns)
    # ----------------------------------------------------
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False

    sig_row = sig_table.rows[0]
    cell_patron = sig_row.cells[0]
    cell_trabajador = sig_row.cells[1]
    cell_patron.width = Inches(3.4)
    cell_trabajador.width = Inches(3.6)

    # Patrón / Dirección
    p_sig1 = cell_patron.paragraphs[0]
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig1.paragraph_format.space_after = Pt(1)
    p_sig1.paragraph_format.space_before = Pt(10)
    r_sig1_hdr = p_sig1.add_run("POR LA EMPRESA / DIRECCIÓN\n\n\n")
    r_sig1_hdr.font.name = 'Segoe UI'
    r_sig1_hdr.font.size = Pt(7.5)
    r_sig1_hdr.font.bold = True
    r_sig1_hdr.font.color.rgb = COLOR_GRIS

    r_sig1_line = p_sig1.add_run("_________________________________________\n")
    r_sig1_line.font.name = 'Segoe UI'
    r_sig1_line.font.color.rgb = COLOR_VERDE

    r_sig1_name = p_sig1.add_run("ING. SOFÍA GAVITO ALARCÓN\n")
    r_sig1_name.font.name = 'Segoe UI'
    r_sig1_name.font.size = Pt(9)
    r_sig1_name.font.bold = True
    r_sig1_name.font.color.rgb = COLOR_CAFE

    r_sig1_cargo = p_sig1.add_run("Dirección General / Titular\nBake Me Up")
    r_sig1_cargo.font.name = 'Segoe UI'
    r_sig1_cargo.font.size = Pt(8)
    r_sig1_cargo.font.color.rgb = COLOR_MUTED

    # Trabajador: José Del Castillo Olvera (Encargado de Producción)
    p_sig2 = cell_trabajador.paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig2.paragraph_format.space_after = Pt(1)
    p_sig2.paragraph_format.space_before = Pt(10)
    r_sig2_hdr = p_sig2.add_run("ACEPTO RESPONSABILIDADES Y RECIBÍ CONFORME\n\n\n")
    r_sig2_hdr.font.name = 'Segoe UI'
    r_sig2_hdr.font.size = Pt(7.5)
    r_sig2_hdr.font.bold = True
    r_sig2_hdr.font.color.rgb = COLOR_GRIS

    r_sig2_line = p_sig2.add_run("_________________________________________\n")
    r_sig2_line.font.name = 'Segoe UI'
    r_sig2_line.font.color.rgb = COLOR_VERDE

    r_sig2_name = p_sig2.add_run("C. JOSÉ DEL CASTILLO OLVERA\n")
    r_sig2_name.font.name = 'Segoe UI'
    r_sig2_name.font.size = Pt(9)
    r_sig2_name.font.bold = True
    r_sig2_name.font.color.rgb = COLOR_CAFE

    r_sig2_cargo = p_sig2.add_run("Encargado de Producción · Bake Me Up\nFecha de acuse: 14 / 10 / 2026")
    r_sig2_cargo.font.name = 'Segoe UI'
    r_sig2_cargo.font.size = Pt(8)
    r_sig2_cargo.font.color.rgb = COLOR_MUTED

    # ----------------------------------------------------
    # 7. FOOTER SECTION (DIRECCIÓN FISCAL EXACTA SEGÚN CSF)
    # ----------------------------------------------------
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_after = Pt(0)
    
    r_ft_line = footer_p.add_run("―" * 48 + "\n")
    r_ft_line.font.name = 'Segoe UI'
    r_ft_line.font.size = Pt(7)
    r_ft_line.font.color.rgb = COLOR_DORADO

    r_ft1 = footer_p.add_run("Bake Me Up · Domicilio Fiscal: Av. Municipio Libre 1972, Local 8, Col. Ex-Hacienda Mayorazgo, C.P. 72480, Puebla, Pue. (Plaza Pabellón Las Torres)\n")
    r_ft1.font.name = 'Segoe UI'
    r_ft1.font.size = Pt(7)
    r_ft1.font.color.rgb = COLOR_MUTED

    r_ft2 = footer_p.add_run("Sitio Oficial: bakemeup.mx · Repostería Fina & Café de Especialidad")
    r_ft2.font.name = 'Segoe UI'
    r_ft2.font.size = Pt(7)
    r_ft2.font.bold = True
    r_ft2.font.color.rgb = COLOR_VERDE

    # ----------------------------------------------------
    # Guardar en rutas oficiales
    # ----------------------------------------------------
    doc_filename = "BakeMeUp_Carta_Incremento_Salarial_Jose_Del_Castillo_14_Octubre_2026.docx"
    doc_filename_simple = "BakeMeUp_Carta_Incremento_Salarial_14_Octubre_2026.docx"
    
    # 1. Repositorio Local Docs
    os.makedirs('docs', exist_ok=True)
    out_repo = os.path.abspath(os.path.join('docs', doc_filename))
    out_repo_simple = os.path.abspath(os.path.join('docs', doc_filename_simple))
    doc.save(out_repo)
    doc.save(out_repo_simple)
    print(f"Guardado en repo: {out_repo}")

    # 2. Carpeta Operativa Interna de BMU: BAKE ME UP ADMINISTRACION
    bmu_admin_root = None
    shortcut_lnk = r"G:\Mi unidad\BAKE ME UP ADMINISTRACION .lnk"
    if os.path.exists(shortcut_lnk):
        try:
            import win32com.client
            shell = win32com.client.Dispatch("WScript.Shell")
            bmu_admin_root = shell.CreateShortcut(shortcut_lnk).TargetPath
        except Exception:
            pass

    if not bmu_admin_root or not os.path.exists(bmu_admin_root):
        candidate = r"G:\.shortcut-targets-by-id\1b2fZ7hp348TbKpb46Y0jqZqoWvkfxHYi\BAKE ME UP ADMINISTRACION"
        if os.path.exists(candidate):
            bmu_admin_root = candidate
        elif os.path.exists(r"G:\Mi unidad\BAKE ME UP ADMINISTRACION"):
            bmu_admin_root = r"G:\Mi unidad\BAKE ME UP ADMINISTRACION"

    if bmu_admin_root and os.path.exists(bmu_admin_root):
        admin_folders = [d for d in os.listdir(bmu_admin_root) if 'admin' in d.lower()]
        if admin_folders:
            admin_dir = os.path.join(bmu_admin_root, admin_folders[0])
            rh_folders = [d for d in os.listdir(admin_dir) if 'recurs' in d.lower() or 'rh' in d.lower() or 'personal' in d.lower()]
            if rh_folders:
                rh_dir = os.path.join(admin_dir, rh_folders[0])
                plantillas_list = [d for d in os.listdir(rh_dir) if 'plantilla' in d.lower()]
                jose_list = [d for d in os.listdir(rh_dir) if 'jose' in d.lower()]
                
                if plantillas_list:
                    p_file = os.path.join(rh_dir, plantillas_list[0], doc_filename)
                    doc.save(p_file)
                    print(f"Guardado en BAKE ME UP ADMINISTRACION (Plantillas): {p_file}")
                
                if jose_list:
                    j_file = os.path.join(rh_dir, jose_list[0], doc_filename)
                    doc.save(j_file)
                    print(f"Guardado en BAKE ME UP ADMINISTRACION (Expediente José): {j_file}")
    else:
        print("Aviso: No se pudo resolver la ruta de BAKE ME UP ADMINISTRACION.")

    return out_repo

if __name__ == '__main__':
    build_salary_increase_doc()

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(13)
    h.paragraph_format.space_after = Pt(4)
    for run in h.runs:
        run.font.name = 'Segoe UI'
        if level == 1:
            run.font.size = Pt(14.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif level == 2:
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif level == 3:
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(100, 116, 139)
    return h

def add_paragraph_styled(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    if text:
        r = p.add_run(text)
        r.font.name = 'Segoe UI'
        r.font.size = Pt(9.5)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet_styled(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(51, 65, 85)
    return p

def build_document():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_badge = doc.add_paragraph()
    p_badge.paragraph_format.space_after = Pt(2)
    r_badge = p_badge.add_run("MEXADEV · PROYECTOS Y TECNOLOGÍA")
    r_badge.font.name = 'Segoe UI'
    r_badge.font.size = Pt(9)
    r_badge.font.bold = True
    r_badge.font.color.rgb = RGBColor(14, 165, 233)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("Bake Me Up: Especificación Técnica de Arquitectura, Requerimientos y Flujo Operativo")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Documento de Implementación para Ingeniería (Lalo / Carlos Ruiz)")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Cliente / Negocio:", "Bake Me Up (Sofia Gavito Alarcón)"),
        ("Equipo de Desarrollo:", "Carlos Ruiz 'Woody' (Líder Proyecto) / Lalo (Lead Backend & Web)"),
        ("Versión del Documento:", "Versión 2.6 - Catálogo de Proveedores, KDS Cocina y Flujo Operativo"),
        ("Fecha de Actualización:", "14 de octubre de 2026")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 40, 40, 80, 80)
        set_cell_margins(c1, 40, 40, 80, 80)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = 'Segoe UI'
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(71, 85, 105)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = 'Segoe UI'
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 1. Objetivo
    add_header_styled(doc, "1. Objetivo del Proyecto", level=1)
    add_paragraph_styled(doc, "Sustituir por completo las hojas de cálculo de Excel/Exportek por una plataforma web integrada en Bake Me Up. La solución aborda cuatro frentes técnicos clave:")
    add_bullet_styled(doc, "Actualización en cascada de recetas base al modificarse el costo de insumos como huevo o lácteos.", "1. Protección de márgenes: ")
    add_bullet_styled(doc, "Cualquier persona del equipo cotiza en < 60 segundos con exactitud matemática.", "2. Estandarización de precios: ")
    add_bullet_styled(doc, "Separación rigurosa de vistas: Panel administrativo para dueñas y KDS táctil en iPad para cocina 100% libre de papel con datos financieros estrictamente bloqueados por backend.", "3. Flujo operativo de taller: ")
    add_bullet_styled(doc, "Directorio centralizado con botón de enlace directo a WhatsApp accesible para todo el equipo.", "4. Catálogo de proveedores: ")

    # 2. Módulos
    add_header_styled(doc, "2. Especificación Técnica de Módulos", level=1)
    
    add_header_styled(doc, "Módulo A: Matriz Dinámica de Insumos y Recetas (Backend de Costos)", level=2)
    add_bullet_styled(doc, "Insumos base con unidad de medida (kg, lt, pza) y precio de compra vigente.", "Catálogo de materias primas: ")
    add_bullet_styled(doc, "El cambio de precio de un ingrediente recalcula en tiempo real las recetas base.", "Actualización en cascada: ")
    add_bullet_styled(doc, "La fórmula del pan toma por defecto el insumo de mayor costo, blindando el margen sin importar el sabor final del cliente.", "Regla de costeo seguro: ")

    add_header_styled(doc, "Módulo B: Motor Web de Cotización Rápida", level=2)
    add_bullet_styled(doc, "Fecha/hora, personas/porciones, estilo (Vintage, Domo, Clásico, Pisos), bizcocho, relleno y adicionales (flores, obleas, papel, MDF).", "Inputs: ")
    add_bullet_styled(doc, "Gramos de betún = f(porciones, estilo). Tiempo extra de mano de obra = f(complejidad, estilo).", "Cálculo paramétrico automático: ")
    add_bullet_styled(doc, "(Costo Insumos + Mano de Obra + Gastos Extra) / (1 - Margen) = Precio Venta Sugerido.", "Fórmula financiera: ")

    add_header_styled(doc, "Módulo C: Histórico y Búsqueda Visual", level=2)
    add_bullet_styled(doc, "Foto de referencia, porciones, sabores, extras y precio cotizado original.", "Ficha histórica: ")
    add_bullet_styled(doc, "Búsqueda semántica (tags) y búsqueda por fotografía de referencia.", "Buscador: ")
    add_bullet_styled(doc, "Comparativa en tiempo real: cuánto se cobró entonces vs. cuánto cuesta elaborarlo hoy.", "Recálculo histórico: ")

    add_header_styled(doc, "Módulo D: Flujo de Pedidos, Calendario y KDS de Cocina (iPad)", level=2)
    add_bullet_styled(doc, "Visualización completa del costo del pastel, anticipo recibido (monto/fecha/método), saldo pendiente y estatus de cobranza (Liquidado / Pendiente). Control exclusivo para marcar 'Entregado'.", "Vista Admin (Dueñas): ")
    add_bullet_styled(doc, "Dispositivo exclusivo en taller, sesión genérica permanente, 100% libre de papel. CRÍTICO PARA LALO: Los endpoints de API para cocina NO retornan campos de dinero (total_price, deposit_amount, balance_due, payment_status).", "Vista KDS Cocina (iPad): ")
    add_bullet_styled(doc, "1) 'Pendiente' (por hornear/decorar), 2) 'En vitrina / refri' (terminado por pastelero y en frío), 3) 'Entregado' (marcado por dueñas al cobrar y entregar al cliente).", "3 Estados Operativos: ")

    add_header_styled(doc, "Módulo E: Catálogo y Directorio de Proveedores", level=2)
    add_paragraph_styled(doc, "Directorio ágil para gestión de compras y contacto con proveedores de la pastelería:")
    add_bullet_styled(doc, "Nombre comercial / Negocio, contacto o persona de atención, teléfono, categoría de insumos que surte (Lácteos, Huevo, Harinas, Flores, MDF / Toppers, Empaques, etc.) y notas operativas (días de visita, mínimos de compra, tiempos de entrega).", "Campos del Proveedor: ")
    add_bullet_styled(doc, "Cada tarjeta o fila de proveedor incluye un botón con enlace directo a WhatsApp (https://wa.me/52...) con mensaje inicial predeterminado para solicitar pedido en 1 clic.", "Botón 1-Clic a WhatsApp: ")
    add_bullet_styled(doc, "Disponible para todo el equipo (taller, mostrador y administración) para que cualquier persona pueda contactar a un proveedor en caso de urgencia.", "Acceso Global de Consulta: ")
    add_bullet_styled(doc, "Alta, edición y baja de proveedores restringida a administradores. Al dar de baja, se aplica borrado lógico (is_active = false) para no romper la integridad referencial de insumos y compras pasadas.", "Gestión y Borrado Seguro (Soft Delete): ")

    # 3. Flujo End-to-End
    add_header_styled(doc, "3. Flujo Operativo Paso a Paso (Happy Path)", level=1)
    pasos = [
        ("Paso 1: Solicitud y Captura (Mostrador/WhatsApp)", "El cliente solicita cotización con fecha/hora, personas, estilo y foto de referencia. Mostrador carga los datos en el cotizador web."),
        ("Paso 2: Generación Instantánea de Precio", "El motor calcula en < 1 segundo el betún en gramos, horas de mano de obra y costo total con margen protegido. Genera resumen de texto limpio para WhatsApp."),
        ("Paso 3: Confirmación de Pedido y Anticipo (Dueñas)", "El cliente acepta. Se cobra el anticipo (ej. 50%), se registra método de pago, el sistema calcula el saldo pendiente y agenda el pedido en el calendario."),
        ("Paso 4: Notificación al Taller (iPad Cocina)", "El pedido ingresa en tiempo real a la columna 'Pendiente' del iPad, ordenado cronológicamente por entrega. Muestra foto, recetas, porciones y extras (sin montos de dinero)."),
        ("Paso 5: Elaboración y Montaje", "El pastelero hornea, rellena, cubre con el betún exacto calculado y decora siguiendo la ficha técnica."),
        ("Paso 6: Finalización y Refrigeración (Listo en Frío)", "El pastelero toca la tarjeta en el iPad y la cambia a 'En vitrina / refri'. En el panel de las dueñas se actualiza al instante sabiendo que el pastel ya está terminado físicamente."),
        ("Paso 7: Llegada del Cliente y Cobro de Saldo", "El cliente acude a mostrador. La dueña revisa en pantalla el saldo pendiente, cobra el restante y registra la liquidación completa."),
        ("Paso 8: Entrega Final y Archivo", "Únicamente la dueña presiona 'Entregado'. El pedido sale de la lista de pendientes activos y se archiva en el historial permanente.")
    ]
    for p_tit, p_desc in pasos:
        add_bullet_styled(doc, p_desc, f"{p_tit}: ")

    # 4. What If
    add_header_styled(doc, "4. Matriz de Escenarios 'What If' (Casos Borde y Excepciones)", level=1)
    what_ifs = [
        ("What If 1: Cotización no confirmada de inmediato (sube costo de insumos)", "Las cotizaciones tienen vigencia configurable (ej. 7-15 días). Si el cliente regresa después, el sistema alerta 'Cotización vencida' y permite recalcular el precio a los costos del día con 1 clic antes de cobrar el anticipo."),
        ("What If 2: Cambio de fecha u hora de entrega de última hora", "Solo las dueñas pueden modificar fecha/hora en Admin. Al editarse, el pedido se reordena automáticamente en el iPad de cocina según su nueva prioridad horaria."),
        ("What If 3: Diseño totalmente atípico fuera de estilos predefinidos", "El cotizador incluye la opción 'Estilo Personalizado/Especial', que permite ingresar manualmente horas estimadas de mano de obra y dificultad sin alterar las fórmulas base de pan y betún."),
        ("What If 4: El cliente llega a recoger pero no liquida el saldo restante", "La dueña no puede marcar 'Entregado' sin liquidar (o el sistema exige confirmación explícita con alerta de deuda). El pastel permanece en 'En vitrina / refri' y el saldo queda registrado como pendiente."),
        ("What If 5: El cliente cancela el pedido", "Si está 'Pendiente', la dueña cancela y se elimina del iPad de cocina inmediatamente para no gastar insumos. Si ya estaba 'En vitrina / refri', el sistema registra la cancelación y el anticipo queda como no reembolsable."),
        ("What If 6: Pastel sufre daño accidental en refrigeración y requiere retrabajo", "El pastelero puede regresar el estado de 'En vitrina / refri' a 'Pendiente'. La tarjeta se resalta visualmente como 'Prioridad / Retrabajo'."),
        ("What If 7: Pérdida temporal de conexión a internet en la cocina", "REQUERIMIENTO TÉCNICO: La interfaz del iPad debe contar con soporte offline (PWA / caché local) para permitir consultar fichas técnicas y fotos del día aun sin internet, sincronizando cambios al reconectar.")
    ]
    for w_tit, w_desc in what_ifs:
        add_bullet_styled(doc, w_desc, f"{w_tit}: ")

    # 5. Criterios para Lalo
    add_header_styled(doc, "5. Criterios de Aceptación y Entregables para Lalo", level=1)
    criterios = [
        ("Base de Datos Relacional", "Modelos limpios para insumos, recetas, estilos, cotizaciones, pedidos y proveedores (con soft delete)."),
        ("Seguridad y RBAC", "Data masking en Backend: ningún dato monetario viaja a la vista de cocina."),
        ("Motor de Cotización", "Cálculo en cascada y coeficientes por estilo en < 1 segundo."),
        ("KDS iPad Touch-friendly", "Interfaz táctil responsive tipo tarjetas con los 3 estados (Pendiente > Refri > Entregado) y soporte offline."),
        ("Panel Admin", "Calendario con saldos, anticipos, validación de entrega y vigencia de cotizaciones."),
        ("Módulo Proveedores", "Directorio con botón directo a WhatsApp (wa.me) accesible para todo el equipo y gestión administrativa.")
    ]
    for tit, desc in criterios:
        add_bullet_styled(doc, desc, f"{tit}: ")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_fin = doc.add_paragraph()
    r_fin = p_fin.add_run("Documento generado por MexaDev · Proyectos y Tecnología.")
    r_fin.font.size = Pt(8.5)
    r_fin.font.italic = True
    r_fin.font.color.rgb = RGBColor(148, 163, 184)

    output_path = r"G:\Mi unidad\MEXADEV\05_Proyectos_y_Tecnologia\BAKEMEUP\01_Especificacion_y_Propuesta\DOCUMENTO_TECNICO_COTIZADOR_Y_MIGRACION_BAKEMEUP.docx"
    doc.save(output_path)
    print("Version 2.6 docx successfully generated at:", output_path)

if __name__ == "__main__":
    build_document()
